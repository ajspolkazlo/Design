"""
Testy dla tools/report_xlsx.py — warstwa raportowania. Bez sieci: mockuje
_final_liveness_check zamiast robic prawdziwe zapytania HTTP.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).parent.parent))
sys.path.insert(0, str(Path(__file__).parent.parent / "tools"))

import report_xlsx as rx


def _row(verify_payload, on_portal_found=1):
    return {"verify_json": json.dumps(verify_payload), "on_portal_found": on_portal_found}


def test_effective_matches_no_recheck_passthrough():
    payload = {"matches": [{"portal": "otodom", "verdict": "CONFIRMED", "url": "http://x", "reasons": []}]}
    matches = rx.effective_matches(_row(payload), recheck_live=False)
    assert matches[0]["verdict"] == "CONFIRMED"


def test_effective_matches_downgrades_stale_link(monkeypatch):
    # Zadanie 1.4: swieza weryfikacja TUZ PRZED raportem — dopasowanie bylo
    # CONFIRMED w bazie, ale w tej chwili link jest martwy -> REJECTED
    monkeypatch.setattr(rx, "_final_liveness_check", lambda portal, url, timeout=10: False)
    payload = {"matches": [{"portal": "otodom", "verdict": "CONFIRMED", "url": "http://x", "reasons": ["stary powod"]}]}
    matches = rx.effective_matches(_row(payload), recheck_live=True)
    assert matches[0]["verdict"] == "REJECTED"
    assert any("nieaktualne" in r for r in matches[0]["reasons"])


def test_effective_matches_keeps_live_link(monkeypatch):
    monkeypatch.setattr(rx, "_final_liveness_check", lambda portal, url, timeout=10: True)
    payload = {"matches": [{"portal": "otodom", "verdict": "CONFIRMED", "url": "http://x", "reasons": []}]}
    matches = rx.effective_matches(_row(payload), recheck_live=True)
    assert matches[0]["verdict"] == "CONFIRMED"


def test_effective_matches_skips_already_rejected():
    # dopasowania juz REJECTED w bazie nie powinny w ogole wywolywac
    # ponownego zapytania sieciowego (tanie, ale bez sensu)
    payload = {"matches": [{"portal": "otodom", "verdict": "REJECTED", "url": "http://x", "reasons": ["wynajem"]}]}
    matches = rx.effective_matches(_row(payload), recheck_live=True)
    assert matches[0]["verdict"] == "REJECTED"
    assert matches[0]["reasons"] == ["wynajem"]  # niezmienione, bez dopisanego powodu


def test_lead_verdict_all_rejected_is_clean():
    matches = [{"portal": "otodom", "verdict": "REJECTED", "url": "http://x", "reasons": ["wynajem"]}]
    verdict, reasons, links, first_url = rx.lead_verdict_and_reasons(matches, checked=True)
    assert verdict == "CLEAN"
    assert links == "" and first_url == ""


def test_lead_verdict_unchecked_when_not_checked_and_no_matches():
    verdict, *_ = rx.lead_verdict_and_reasons([], checked=False)
    assert verdict == "UNCHECKED"


def test_per_portal_matches_maps_by_portal():
    matches = [
        {"portal": "otodom", "verdict": "CONFIRMED", "url": "http://a"},
        {"portal": "olx", "verdict": "REJECTED", "url": "http://b"},
    ]
    out = rx.per_portal_matches(matches)
    assert out["otodom"] == ("CONFIRMED", "http://a")
    assert out["olx"] == ("REJECTED", "http://b")
    assert "gratka" not in out


# --------------------------- build(): kolumny Strona dewelopera / Status strony ---------------------------

def _fake_leads_df():
    base = {
        "score": 50, "data": "2026-01-01", "gmina": "Piaseczno", "miejscowosc": "Piaseczno",
        "adres_pelny": "ul. Testowa 1, Piaseczno", "kategoria_obiektu": "budynek wielorodzinny",
        "distance_km": 10.0, "lat": 52.0, "lon": 21.0,
        "verify_json": json.dumps({"matches": [], "investor_serial_count": 1,
                                    "parcel_owner_desc": "b/d", "parcel_area_m2": None}),
        "on_portal_found": 1,
    }
    rows = [
        {**base, "id_sprawy": "ST-1", "inwestor": "TOP INVESTMENT Sp. z o.o.",
         "dev_site_url": "https://topinvestment.pl/", "dev_site_status": "potwierdzona"},
        {**base, "id_sprawy": "ST-2", "inwestor": "Nieznana Firma Sp. z o.o.",
         "dev_site_url": "https://companiesmarketcap.com/x", "dev_site_status": "kandydat_niepewny"},
        {**base, "id_sprawy": "ST-3", "inwestor": "Jan Kowalski",
         "dev_site_url": np.nan, "dev_site_status": "brak_do_wyszukania_osoba_fizyczna"},
    ]
    return pd.DataFrame(rows)


def test_build_writes_dev_site_columns(monkeypatch, tmp_path):
    monkeypatch.setattr(rx, "load_rows", lambda db_path, status: _fake_leads_df())
    out_path = tmp_path / "raport.xlsx"
    rx.build(out_path, status="scored", recheck_live=False)

    from openpyxl import load_workbook
    wb = load_workbook(out_path)
    ws = wb["Leady"]
    headers = [c.value for c in ws[1]]
    assert "Strona dewelopera" in headers
    assert "Status strony" in headers
    site_col = headers.index("Strona dewelopera") + 1
    status_col = headers.index("Status strony") + 1

    # wiersz 1 (ST-1): potwierdzona -> klikalny hiperlacz
    c1 = ws.cell(row=2, column=site_col)
    assert c1.value == "https://topinvestment.pl/"
    assert c1.hyperlink is not None
    assert ws.cell(row=2, column=status_col).value == "✔ potwierdzona"

    # wiersz 2 (ST-2): kandydat_niepewny -> URL widoczny jako tekst, ALE bez hiperlacza
    c2 = ws.cell(row=3, column=site_col)
    assert c2.value == "https://companiesmarketcap.com/x"
    assert c2.hyperlink is None
    assert "niepewna" in ws.cell(row=3, column=status_col).value

    # wiersz 3 (ST-3): osoba fizyczna -> brak URL-a w ogole
    c3 = ws.cell(row=4, column=site_col)
    assert c3.value == "—"
    assert c3.hyperlink is None
